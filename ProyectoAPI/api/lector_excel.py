from openpyxl import load_workbook

# Columnas que necesita el proyecto: nombre en el Excel -> nombre corto que usaremos en el código
COLUMNAS = {
    "Departamento": "departamento",
    "Municipio": "municipio",
    "Cultivo": "cultivo",
    "Topografia": "topografia",
    "pH agua:suelo 2,5:1,0": "ph",
    "Fósforo (P) Bray II mg/kg": "fosforo",
    "Potasio (K) intercambiable cmol(+)/kg": "potasio",
}


def cargar_registros(ruta_excel):
    libro = load_workbook(ruta_excel, read_only=True, data_only=True)
    hoja = libro.active

    filas = hoja.iter_rows(values_only=True)
    encabezados = [str(h).strip() if h is not None else "" for h in next(filas)]

    # posición de cada columna que nos interesa
    posiciones = {}
    for nombre_excel, nombre_corto in COLUMNAS.items():
        posiciones[nombre_corto] = encabezados.index(nombre_excel)

    registros = []
    for fila in filas:
        registro = {}
        for nombre_corto, pos in posiciones.items():
            registro[nombre_corto] = fila[pos]
        if registro["departamento"] is None:  # fila vacía al final del Excel
            continue
        registros.append(registro)

    libro.close()
    return registros